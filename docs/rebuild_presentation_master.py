import sys, os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

sys.stdout.reconfigure(encoding='utf-8')

src = r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh_khoi_phuc_goc.pptx'
dst = r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh.pptx'

prs = Presentation(src)

COLOR_DARK_BLUE = RGBColor(11, 44, 88)
COLOR_GOLD = RGBColor(218, 142, 17)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_TEXT_DARK = RGBColor(33, 37, 41)
COLOR_TEXT_MUTED = RGBColor(108, 117, 125)

def hide_shape(sh):
    sh.left = -10000000
    sh.top = -10000000
    if sh.has_text_frame:
        sh.text_frame.text = ""

def set_text(sh, text, font_size=Pt(10), bold=False, color=COLOR_TEXT_DARK, align=PP_ALIGN.LEFT):
    if not sh.has_text_frame:
        return
    sh.text_frame.word_wrap = True
    sh.text_frame.text = text
    for p in sh.text_frame.paragraphs:
        p.font.name = "Calibri"
        p.font.size = font_size
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align

# ==========================================
# SLIDE 1: TRANG BÌA
# ==========================================
s1 = prs.slides[0]
for sh in s1.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0':
            set_text(sh, "TRƯỜNG ĐẠI HỌC THỦ DẦU MỘT · VIỆN CÔNG NGHỆ SỐ", Pt(11), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif name == 'Text 2': # Dark blue pill
            sh.left = 3504682
            sh.width = 5169481
            set_text(sh, "BÁO CÁO ĐỒ ÁN CƠ SỞ NGÀNH", Pt(13), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 3':
            set_text(sh, "Xây dựng Website Truyền thông Công đoàn TDMU\ntích hợp AI hỗ trợ tạo và đăng tải nội dung", Pt(24), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif name == 'Text 4':
            sh.left = 2500000
            sh.width = 7200000
            set_text(sh, "GVHD: ThS. Võ Quốc Lương   |   Nhóm thực hiện: Nhóm 2", Pt(12), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif name == 'Text 5' and sh.shape_id == 2:
            sh.left = 2500000
            sh.width = 7200000
            sh.top = 4500000
            sh.height = 900000
            members_text = (
                "1. Nguyễn Bình Dương (Trưởng nhóm) - MSSV: 2424802010319 - Lớp: D24CNTT05\n"
                "2. Trần Hồng Thanh - MSSV: 2424802010439 - Lớp: D24CNTT03\n"
                "3. Phạm Anh Tuấn - MSSV: 2324802010393 - Lớp: D23CNTT03"
            )
            set_text(sh, members_text, Pt(11), False, COLOR_TEXT_DARK, PP_ALIGN.CENTER)
        elif name == 'Text 5' and sh.shape_id != 2:
            hide_shape(sh)
        elif name == 'Text 6':
            set_text(sh, "HỌC PHẦN: ĐỒ ÁN CƠ SỞ NGÀNH · NĂM HỌC 2026–2027 · THÁNG 10/2026", Pt(9.5), False, COLOR_TEXT_MUTED, PP_ALIGN.CENTER)

# ==========================================
# SLIDE 2: BỐI CẢNH & TÍNH CẤP THIẾT
# ==========================================
s2 = prs.slides[1]
for sh in s2.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0':
            set_text(sh, "1. GIỚI THIỆU ĐỀ TÀI", Pt(11), True, COLOR_GOLD)
        elif name == 'Text 1':
            set_text(sh, "Bối cảnh & Tính cấp thiết của đề tài", Pt(18), True, COLOR_DARK_BLUE)
        elif name == 'Text 3':
            set_text(sh, "Công đoàn Trường ĐH Thủ Dầu Một", Pt(14), True, COLOR_DARK_BLUE)
        elif name == 'Text 4':
            desc_left = (
                "• Quy mô quản lý: Hơn 800 cán bộ, giảng viên và người lao động trong toàn trường.\n\n"
                "• Kênh truyền thông hiện tại: Mục tin trên Website trường, Fanpage Facebook, các nhóm Zalo phân tán.\n\n"
                "• Thực trạng nhân lực: Đội ngũ làm truyền thông là giảng viên kiêm nhiệm, khối lượng công việc chuyên môn lớn, thời gian dành cho tin bài hạn chế."
            )
            set_text(sh, desc_left, Pt(10.5), False, COLOR_TEXT_DARK)
            sh.top = 1800000
            sh.height = 2200000
        elif name in ['Text 7', 'Text 9', 'Text 10', 'Text 11']:
            hide_shape(sh)
        elif name == 'Text 14':
            set_text(sh, "01", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 15':
            set_text(sh, "Áp lực kiêm nhiệm & Chậm trễ tin tức", Pt(12), True, COLOR_DARK_BLUE)
        elif name == 'Text 16':
            c1_desc = "Cán bộ truyền thông kiêm nhiệm nhiều nhiệm vụ đào tạo; khâu chuẩn bị tư liệu, chấp bút tin bài theo chuẩn phong cách báo chí đoàn thể tốn nhiều thời gian, dễ gây chậm trễ thông tin."
            set_text(sh, c1_desc, Pt(10), False, COLOR_TEXT_DARK)
        elif name == 'Text 19':
            set_text(sh, "02", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 20':
            set_text(sh, "Đăng tải thủ công, phân tán nhiều kênh", Pt(12), True, COLOR_DARK_BLUE)
        elif name == 'Text 21':
            c2_desc = "Quy trình xuất bản rời rạc; cán bộ phải sao chép văn bản, tự biên tập lại ảnh và đăng tải lặp đi lặp lại thủ công trên từng kênh (Website, Facebook, Zalo), dễ sai sót thông tin."
            set_text(sh, c2_desc, Pt(10), False, COLOR_TEXT_DARK)
        elif name == 'Text 24':
            set_text(sh, "03", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 25':
            set_text(sh, "Thiếu công cụ đo lường tập trung", Pt(12), True, COLOR_DARK_BLUE)
        elif name == 'Text 26':
            c3_desc = "Chưa có công cụ định lượng tập trung (lượt xem, tương tác đa kênh) để Ban Chấp hành kịp thời đánh giá và tối ưu hiệu quả truyền thông."
            sh.height = 360000
            set_text(sh, c3_desc, Pt(9.5), False, COLOR_TEXT_DARK)
        elif name == 'Text 38':
            set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# ==========================================
# SLIDE 3: MỤC TIÊU & GIẢI PHÁP
# ==========================================
s3 = prs.slides[2]
for sh in s3.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0':
            set_text(sh, "1. GIỚI THIỆU ĐỀ TÀI", Pt(11), True, COLOR_GOLD)
        elif name == 'Text 1':
            set_text(sh, "Mục tiêu nghiên cứu & Định hướng giải quyết", Pt(18), True, COLOR_DARK_BLUE)
        elif name == 'Text 8':
            set_text(sh, "01", Pt(13), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif name == 'Text 9':
            set_text(sh, "Mục tiêu trọng tâm của đề tài", Pt(13), True, COLOR_DARK_BLUE)
        elif name == 'Text 10':
            t10 = "Xây dựng hệ thống website truyền thông tích hợp Trợ lý AI hỗ trợ tạo bài viết từ tư liệu sự kiện và tự động hóa quy trình đăng tải đa kênh (Website Công đoàn, Fanpage Facebook, Zalo OA), giải phóng sức lao động thủ công và nâng cao hiệu quả thông tin đoàn thể."
            set_text(sh, t10, Pt(10.5), False, COLOR_TEXT_DARK)
        elif name == 'Text 17':
            set_text(sh, "02", Pt(13), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif name == 'Text 18':
            set_text(sh, "Hỗ trợ tạo nội dung bằng AI", Pt(13), True, COLOR_DARK_BLUE)
        elif name == 'Text 19':
            ai_sol = (
                "• Tự động phân tích kế hoạch, thông báo sự kiện để lập dàn ý và dựng bài 5W1H.\n"
                "• Chấp bút đoạn mở đầu (Sapo) thu hút, gợi ý 3–5 tiêu đề báo chí chuẩn mực.\n"
                "• Tự động rà soát lỗi chính tả tiếng Việt, chuẩn hóa văn phong hành chính đoàn thể.\n"
                "• Chuyển đổi linh hoạt định dạng tin tức dài sang bài đăng ngắn kèm hashtag cho MXH."
            )
            set_text(sh, ai_sol, Pt(10), False, COLOR_TEXT_DARK)
        elif name == 'Text 26':
            set_text(sh, "03", Pt(13), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif name == 'Text 27':
            set_text(sh, "Xuất bản đa kênh tự động & Đo lường", Pt(13), True, COLOR_DARK_BLUE)
        elif name == 'Text 28':
            pub_sol = (
                "• Tự động đăng tải nội dung và hình ảnh đồng thời lên Website, Facebook Fanpage.\n"
                "• Cơ chế lập lịch hẹn giờ xuất bản ngầm (Background Scheduler) chính xác, tin cậy.\n"
                "• Phân định 4 vai trò rõ rệt: Quản trị viên, Biên tập viên, Cộng tác viên, Độc giả.\n"
                "• Bảng điều khiển (Dashboard) theo dõi định lượng lượt xem và độ lan tỏa tin tức."
            )
            set_text(sh, pub_sol, Pt(10), False, COLOR_TEXT_DARK)
        elif name == 'Text 35':
            set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# ==========================================
# SLIDE 4: CỔNG THÔNG TIN ĐỘC GIẢ (Mapped đúng từng shape)
# ==========================================
s4 = prs.slides[3]
c1 = (
    "• Phân loại bài viết rõ ràng theo 8 chuyên mục công đoàn đặc thù.\n"
    "• Tìm kiếm toàn văn thông minh theo từ khóa, chuyên mục, tác giả và ngày đăng.\n"
    "• Phân trang tối ưu, hiển thị tin nổi bật trên Slider trang chủ tạo điểm nhấn."
)
c2 = (
    "• Trình bày bài viết theo bố cục báo điện tử: Tiêu đề, Sapo nổi bật, Thân bài rich-text.\n"
    "• Album ảnh sự kiện độ phân giải cao kèm chú thích chân ảnh chi tiết.\n"
    "• Tệp văn bản đính kèm cho phép xem trước trực tuyến và tải về an toàn."
)
c3 = (
    "• Đoàn viên bày tỏ cảm xúc (Like), bình luận trao đổi văn minh dưới bài viết.\n"
    "• Chia sẻ một chạm liên kết bài viết sang Facebook, Zalo cá nhân nhanh chóng.\n"
    "• Hộp thư điện tử tiếp nhận ý kiến đóng góp, nguyện vọng của người lao động."
)
c4 = (
    "• Tra cứu hệ thống nghị quyết, quyết định, quy chế công đoàn minh bạch.\n"
    "• Tải về biểu mẫu xin trợ cấp khó khăn, mẫu đề xuất khen thưởng dạng Word/PDF.\n"
    "• Phân loại văn bản theo năm ban hành và cơ quan thẩm quyền rõ ràng."
)

for sh in s4.shapes:
    name = sh.name
    if name == 'Text 0':
        set_text(sh, "2. PHÂN TÍCH YÊU CẦU", Pt(11), True, COLOR_GOLD)
    elif name == 'Text 1':
        set_text(sh, "Cổng thông tin độc giả — Đọc báo, tra cứu & tương tác", Pt(18), True, COLOR_DARK_BLUE)
    elif name in ['Text 3', 'Text 7', 'Text 12', 'Text 17']:
        if sh.has_text_frame:
            sh.text_frame.text = "" # Clean background container shapes
    elif name == 'Text 4': # Badge 01
        set_text(sh, "01", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif name == 'Text 5': # Title 01
        set_text(sh, "Tin tức phong trào & Phân loại đa tầng", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 6': # Desc 01
        sh.width = 4600000; sh.height = 1300000
        set_text(sh, c1, Pt(10), False, COLOR_TEXT_DARK)

    elif name == 'Text 9': # Badge 02
        set_text(sh, "02", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif name == 'Text 10': # Title 02
        set_text(sh, "Đọc bài viết đa phương tiện chuẩn báo chí", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 11': # Desc 02
        sh.width = 4600000; sh.height = 1300000
        set_text(sh, c2, Pt(10), False, COLOR_TEXT_DARK)

    elif name == 'Text 14': # Badge 03
        set_text(sh, "03", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif name == 'Text 15': # Title 03
        set_text(sh, "Tương tác độc giả & Chia sẻ liên kết", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 16': # Desc 03
        sh.width = 4600000; sh.height = 1300000
        set_text(sh, c3, Pt(10), False, COLOR_TEXT_DARK)

    elif name == 'Text 19': # Badge 04
        set_text(sh, "04", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif name == 'Text 20': # Title 04
        set_text(sh, "Kho văn bản pháp quy & Biểu mẫu điện tử", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 21': # Desc 04
        sh.width = 4600000; sh.height = 1300000
        set_text(sh, c4, Pt(10), False, COLOR_TEXT_DARK)
    elif name == 'Text 33':
        set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# ==========================================
# SLIDE 5: PHÂN HỆ CMS
# ==========================================
s5 = prs.slides[4]
s5_data = [
    ("01", "Phân quyền 4 tác nhân", "Cấu hình vai trò rõ rệt: Admin quản trị, Editor biên tập, Contributor soạn bài, User đọc báo."),
    ("02", "Soạn thảo trực quan", "Trình biên tập đa phương tiện phong phú, chèn bảng biểu, định dạng báo chí, tải album ảnh sự kiện."),
    ("03", "Quản lý trạng thái bài", "Quy trình mạch lạc: Bản nháp → Đã gửi → Xuất bản ngay / Hẹn giờ phát hành; xem trước đa kênh."),
    ("04", "Lập lịch tự động", "Thiết lập thời gian đăng bài linh hoạt; tiến trình ngầm (Node-cron) tự động kích hoạt xuất bản đúng giờ."),
    ("05", "Đồng bộ Fanpage Facebook", "Tích hợp Meta Graph API v19 đăng tự động tóm tắt, ảnh đại diện và liên kết về cổng thông tin."),
    ("06", "Bảng điều hành (Dashboard)", "Theo dõi trực quan tổng số tin bài, trạng thái xuất bản và biểu đồ định lượng mức độ tương tác."),
    ("07", "Quản lý Văn bản & Mẫu biểu", "Cập nhật số hiệu văn bản, trích yếu, tệp scan PDF và biểu mẫu hành chính tải về theo danh mục."),
    ("08", "Quản trị hệ thống & AI Token", "Quản lý tài khoản người dùng, phân quyền RBAC, quản lý khóa API Gemini và hạn mức Token an toàn.")
]
for sh in s5.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0':
            set_text(sh, "2. PHÂN TÍCH YÊU CẦU", Pt(11), True, COLOR_GOLD)
        elif name == 'Text 1':
            set_text(sh, "Phân hệ Quản trị CMS & Điều phối xuất bản đa kênh", Pt(18), True, COLOR_DARK_BLUE)
        elif name == 'Text 38':
            set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

card_num_names = ['Text 5', 'Text 9', 'Text 13', 'Text 17', 'Text 21', 'Text 25', 'Text 29', 'Text 33']
card_title_names = ['Text 6', 'Text 10', 'Text 14', 'Text 18', 'Text 22', 'Text 26', 'Text 30', 'Text 34']
card_desc_names = ['Text 7', 'Text 11', 'Text 15', 'Text 19', 'Text 23', 'Text 27', 'Text 31', 'Text 35']

for i in range(8):
    for sh in s5.shapes:
        if sh.name == card_num_names[i]:
            set_text(sh, s5_data[i][0], Pt(11), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif sh.name == card_title_names[i]:
            set_text(sh, s5_data[i][1], Pt(11), True, COLOR_DARK_BLUE)
        elif sh.name == card_desc_names[i]:
            set_text(sh, s5_data[i][2], Pt(9.5), False, COLOR_TEXT_DARK)

# ==========================================
# SLIDE 6: TRỢ LÝ AI (2x2 Grid cân đối)
# ==========================================
s6 = prs.slides[5]
for sh in s6.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0':
            set_text(sh, "2. PHÂN TÍCH YÊU CẦU", Pt(11), True, COLOR_GOLD)
        elif name == 'Text 1':
            set_text(sh, "Trợ lý AI hỗ trợ sáng tạo nội dung đa kênh", Pt(18), True, COLOR_DARK_BLUE)
        elif name == 'Text 2':
            set_text(sh, "Ứng dụng Google Gemini AI đọc hiểu tư liệu thô, tự động tạo bài viết và tối ưu định dạng phát hành theo từng kênh.", Pt(10.5), False, COLOR_TEXT_DARK)
        elif name == 'Text 34':
            set_text(sh, "Hạ tầng kỹ thuật: Tích hợp Google Gemini API qua giao thức an toàn, xử lý dữ liệu thời gian thực và quản lý hạn mức Token.", Pt(10.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 36':
            set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# Card 1:
for sh in s6.shapes:
    if sh.name == 'Text 3':
        sh.left = 533400; sh.top = 1573709; sh.width = 5400000; sh.height = 1800000
    elif sh.name == 'Text 4':
        sh.left = 714375; sh.top = 1726109; sh.width = 310000; sh.height = 310000
    elif sh.name == 'Text 5':
        sh.left = 714375; sh.top = 1726109; sh.width = 310000; sh.height = 310000
        set_text(sh, "01", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif sh.name == 'Text 6':
        sh.left = 1150000; sh.top = 1726109; sh.width = 4600000
        set_text(sh, "Chấp bút bài viết từ tư liệu sự kiện", Pt(12), True, COLOR_DARK_BLUE)
    elif sh.name == 'Text 7':
        sh.left = 714375; sh.top = 2107109; sh.width = 5000000; sh.height = 1100000
        t = "Đọc hiểu kế hoạch, thông báo, ảnh sự kiện (DOCX, PDF, ảnh scan) để tự động dựng bản thảo bài viết hoàn chỉnh theo cấu trúc chuẩn 5W1H của báo chí đoàn thể."
        set_text(sh, t, Pt(10), False, COLOR_TEXT_DARK)

# Card 2:
for sh in s6.shapes:
    if sh.name == 'Text 8':
        sh.left = 6250000; sh.top = 1573709; sh.width = 5400000; sh.height = 1800000
    elif sh.name == 'Text 9':
        sh.left = 6430975; sh.top = 1726109; sh.width = 310000; sh.height = 310000
    elif sh.name == 'Text 10':
        sh.left = 6430975; sh.top = 1726109; sh.width = 310000; sh.height = 310000
        set_text(sh, "02", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif sh.name == 'Text 11':
        sh.left = 6866600; sh.top = 1726109; sh.width = 4600000
        set_text(sh, "Tóm tắt Sapo & Định dạng đa kênh", Pt(12), True, COLOR_DARK_BLUE)
    elif sh.name == 'Text 12':
        sh.left = 6430975; sh.top = 2107109; sh.width = 5000000; sh.height = 1100000
        t = "Tự động tạo đoạn mở đầu (Sapo) cô đọng cho Website; đồng thời rút gọn thành các đoạn tóm tắt súc tích kèm hashtag phù hợp cho bài đăng Facebook và Zalo."
        set_text(sh, t, Pt(10), False, COLOR_TEXT_DARK)

# Card 3:
for sh in s6.shapes:
    if sh.name == 'Text 23':
        sh.left = 533400; sh.top = 3550000; sh.width = 5400000; sh.height = 1800000
    elif sh.name == 'Text 24':
        sh.left = 714375; sh.top = 3702400; sh.width = 310000; sh.height = 310000
    elif sh.name == 'Text 25':
        sh.left = 714375; sh.top = 3702400; sh.width = 310000; sh.height = 310000
        set_text(sh, "03", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif sh.name == 'Text 26':
        sh.left = 1150000; sh.top = 3702400; sh.width = 4600000
        set_text(sh, "Gợi ý tiêu đề bài viết chuẩn SEO", Pt(12), True, COLOR_DARK_BLUE)
    elif sh.name == 'Text 27':
        sh.left = 714375; sh.top = 4080000; sh.width = 5000000; sh.height = 1100000
        t = "Phân tích nội dung cốt lõi của bài báo để đề xuất 3–5 phương án tiêu đề trang trọng, đúng văn phong công đoàn nhưng vẫn thu hút độc giả và tối ưu công cụ tìm kiếm."
        set_text(sh, t, Pt(10), False, COLOR_TEXT_DARK)

# Card 4:
for sh in s6.shapes:
    if sh.name == 'Text 28':
        sh.left = 6250000; sh.top = 3550000; sh.width = 5400000; sh.height = 1800000
    elif sh.name == 'Text 29':
        sh.left = 6430975; sh.top = 3702400; sh.width = 310000; sh.height = 310000
    elif sh.name == 'Text 30':
        sh.left = 6430975; sh.top = 3702400; sh.width = 310000; sh.height = 310000
        set_text(sh, "04", Pt(12), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif sh.name == 'Text 31':
        sh.left = 6866600; sh.top = 3702400; sh.width = 4600000
        set_text(sh, "Chuẩn hóa văn phong & Ngữ pháp", Pt(12), True, COLOR_DARK_BLUE)
    elif sh.name == 'Text 32':
        sh.left = 6430975; sh.top = 4080000; sh.width = 5000000; sh.height = 1100000
        t = "Rà soát chính tả tiếng Việt, tinh chỉnh câu từ mạch lạc, thay thế từ ngữ khẩu ngữ bằng thuật ngữ chuẩn mực theo phong cách văn bản hành chính Công đoàn Việt Nam."
        set_text(sh, t, Pt(10), False, COLOR_TEXT_DARK)

# ==========================================
# SLIDE 7: KIẾN TRÚC CÔNG NGHỆ
# ==========================================
s7 = prs.slides[6]
for sh in s7.shapes:
    name = sh.name
    if name == 'Text 0':
        set_text(sh, "3. GIẢI PHÁP CÔNG NGHỆ", Pt(11), True, COLOR_GOLD)
    elif name == 'Text 1':
        set_text(sh, "Kiến trúc kỹ thuật — 4 thành phần hệ thống", Pt(18), True, COLOR_DARK_BLUE)
    elif name == 'Text 7':
        set_text(sh, "Frontend", Pt(14), True, COLOR_DARK_BLUE)
    elif name == 'Text 8':
        set_text(sh, "GIAO DIỆN NGƯỜI DÙNG", Pt(9.5), True, COLOR_TEXT_MUTED, PP_ALIGN.RIGHT)
    elif name == 'Text 11':
        sh.left = 771525; sh.top = 1820000; sh.width = 3200000
        set_text(sh, "React 18 SPA · Tailwind CSS", Pt(10), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
    elif name == 'Text 14':
        sh.top = 2250000; sh.height = 1000000
        set_text(sh, "Cổng thông tin và CMS cơ chế tải nhanh Single Page Application, chuẩn nhận diện thương hiệu TDMU, tối ưu hiển thị Responsive trên PC và thiết bị di động.", Pt(10), False, COLOR_TEXT_DARK)

    elif name == 'Text 20':
        set_text(sh, "Backend", Pt(14), True, COLOR_DARK_BLUE)
    elif name == 'Text 21':
        set_text(sh, "XỬ LÝ NGHIỆP VỤ", Pt(9.5), True, COLOR_TEXT_MUTED, PP_ALIGN.RIGHT)
    elif name == 'Text 22':
        sh.left = 6429375; sh.top = 1820000; sh.width = 3400000
        set_text(sh, "Node.js · Express RESTful API", Pt(10), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
    elif name in ['Text 24', 'Text 28']:
        hide_shape(sh)
    elif name == 'Text 29':
        sh.top = 2250000; sh.height = 1000000
        set_text(sh, "Xử lý nghiệp vụ tập trung, xác thực bảo mật JWT, mã hóa Bcrypt và điều phối tiến trình ngầm (Node-cron) tự động lập lịch xuất bản tin bài đúng giờ.", Pt(10), False, COLOR_TEXT_DARK)

    elif name == 'Text 35':
        set_text(sh, "Cơ sở dữ liệu", Pt(14), True, COLOR_DARK_BLUE)
    elif name == 'Text 36':
        set_text(sh, "LƯU TRỮ & TOÀN VẸN", Pt(9.5), True, COLOR_TEXT_MUTED, PP_ALIGN.RIGHT)
    elif name == 'Text 37':
        sh.left = 771524; sh.top = 4180000; sh.width = 3200000
        set_text(sh, "Microsoft SQL Server 2022", Pt(10), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
    elif name == 'Text 42':
        sh.top = 4600000; sh.height = 1000000
        set_text(sh, "Mô hình dữ liệu quan hệ chuẩn hóa 3NF, bảo đảm toàn vẹn giao dịch ACID, tối ưu hóa truy vấn thông qua hệ thống chỉ mục (Indexing) và quan hệ chặt chẽ.", Pt(10), False, COLOR_TEXT_DARK)

    elif name == 'Text 48':
        set_text(sh, "Dịch vụ tích hợp", Pt(14), True, COLOR_DARK_BLUE)
    elif name == 'Text 49':
        set_text(sh, "AI & XUẤT BẢN ĐA KÊNH", Pt(9.5), True, COLOR_TEXT_MUTED, PP_ALIGN.RIGHT)
    elif name == 'Text 50':
        sh.left = 6429375; sh.top = 4180000; sh.width = 2500000
        set_text(sh, "Google Gemini API", Pt(10), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
    elif name == 'Text 53':
        sh.left = 9100000; sh.top = 4180000; sh.width = 2400000
        set_text(sh, "Meta Graph API v19", Pt(10), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
    elif name == 'Text 54':
        sh.top = 4600000; sh.height = 1000000
        set_text(sh, "Gemini xử lý ngôn ngữ tự nhiên hỗ trợ tạo bài viết; Meta Graph API cung cấp kết nối an toàn tự động đăng bài lên Fanpage ngay khi có lệnh xuất bản.", Pt(10), False, COLOR_TEXT_DARK)
    elif name == 'Text 56':
        set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# ==========================================
# SLIDE 8: USE CASE VÀ 4 TÁC NHÂN
# ==========================================
s8 = prs.slides[7]
for sh in s8.shapes:
    name = sh.name
    if name == 'Text 0':
        set_text(sh, "4. PHÂN TÍCH & THIẾT KẾ HỆ THỐNG", Pt(11), True, COLOR_GOLD)
    elif name == 'Text 1':
        set_text(sh, "Biểu đồ Use Case tổng quát và 4 nhóm tác nhân", Pt(18), True, COLOR_DARK_BLUE)
    elif name == 'Text 2':
        set_text(sh, "4 NHÓM TÁC NHÂN (ACTOR)", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 3':
        sh.left = 533400; sh.top = 1450000; sh.width = 1350000; sh.height = 340000
        set_text(sh, "Quản trị viên\n(Admin)", Pt(9.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif name == 'Text 4':
        sh.left = 1980000; sh.top = 1450000; sh.width = 1350000; sh.height = 340000
        set_text(sh, "Biên tập viên\n(Editor)", Pt(9.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif name == 'Text 5':
        sh.left = 3430000; sh.top = 1450000; sh.width = 1450000; sh.height = 340000
        set_text(sh, "Cộng tác viên\n(Contributor)", Pt(9.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif name == 'Text 6':
        sh.left = 4980000; sh.top = 1450000; sh.width = 1450000; sh.height = 340000
        set_text(sh, "Đoàn viên / Độc giả\n(User)", Pt(9.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
    elif name == 'Text 7':
        sh.top = 1920000
        set_text(sh, "5 PHÂN HỆ NGHIỆP VỤ (26 USE CASE)", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 8':
        sh.top = 2200000; sh.height = 2400000
    elif name == 'Text 10':
        sh.top = 2300000; set_text(sh, "1", Pt(11), True, COLOR_GOLD, PP_ALIGN.CENTER)
    elif name == 'Text 11':
        sh.top = 2300000; set_text(sh, "1. Quản lý bài viết & nội dung — Soạn thảo, đính kèm media, lưu nháp.", Pt(10), False, COLOR_TEXT_DARK)
    elif name == 'Text 13':
        sh.top = 2700000; set_text(sh, "2", Pt(11), True, COLOR_GOLD, PP_ALIGN.CENTER)
    elif name == 'Text 14':
        sh.top = 2700000; set_text(sh, "2. Trợ lý AI hỗ trợ sáng tạo — Dựng bài 5W1H, gợi ý tiêu đề, tóm tắt Sapo.", Pt(10), False, COLOR_TEXT_DARK)
    elif name == 'Text 16':
        sh.top = 3100000; set_text(sh, "3", Pt(11), True, COLOR_GOLD, PP_ALIGN.CENTER)
    elif name == 'Text 17':
        sh.top = 3100000; set_text(sh, "3. Xuất bản đa kênh & Lập lịch — Xem trước, hẹn giờ, đồng bộ Facebook.", Pt(10), False, COLOR_TEXT_DARK)
    elif name == 'Text 19':
        sh.top = 3500000; set_text(sh, "4", Pt(11), True, COLOR_GOLD, PP_ALIGN.CENTER)
    elif name == 'Text 23':
        sh.top = 3500000; set_text(sh, "4. Cổng thông tin độc giả — Đọc tin tức, tra cứu văn bản pháp quy, gửi góp ý.", Pt(10), False, COLOR_TEXT_DARK)
    elif name == 'Text 22':
        sh.top = 3900000; set_text(sh, "5", Pt(11), True, COLOR_GOLD, PP_ALIGN.CENTER)
    elif name == 'Text 25':
        sh.top = 3900000; set_text(sh, "5. Quản trị hệ thống & Người dùng — Quản lý tài khoản, phân quyền RBAC.", Pt(10), False, COLOR_TEXT_DARK)
    elif name == 'Text 26':
        sh.top = 4750000; sh.height = 680000
    elif name == 'Text 27':
        sh.top = 4850000; sh.height = 550000
        note = (
            "• Phân định vai trò: Admin quản trị hệ thống; Editor quản lý tin & duyệt bài; Contributor soạn & gửi bài; User đọc & tương tác.\n"
            "• Toàn bộ 26 use case được đặc tả chuẩn UML với đầy đủ quan hệ <<include>>, <<extend>>."
        )
        set_text(sh, note, Pt(9.5), False, COLOR_TEXT_DARK)
    elif name == 'Text 29':
        set_text(sh, "Biểu đồ Use Case tổng quát hệ thống", Pt(10), True, COLOR_TEXT_MUTED, PP_ALIGN.CENTER)
    elif name == 'Text 31':
        set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# ==========================================
# SLIDE 9: CƠ SỞ DỮ LIỆU
# ==========================================
s9 = prs.slides[8]
for sh in s9.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0':
            set_text(sh, "4. PHÂN TÍCH & THIẾT KẾ HỆ THỐNG", Pt(11), True, COLOR_GOLD)
        elif name == 'Text 1':
            set_text(sh, "Thiết kế cơ sở dữ liệu quan hệ (15 bảng chuẩn 3NF)", Pt(18), True, COLOR_DARK_BLUE)
        elif name == 'Text 3':
            set_text(sh, "Đặc tả kỹ thuật CSDL", Pt(12), True, COLOR_DARK_BLUE)
        elif name == 'Text 4':
            c_db = (
                "• Hệ quản trị: Microsoft SQL Server 2022 kết hợp Connection Pool tối ưu hiệu năng.\n"
                "• Chuẩn hóa dữ liệu: Thiết kế chuẩn bậc 3 (3NF), loại bỏ dư thừa và hạn chế dị thường.\n"
                "• Ràng buộc toàn vẹn: Khóa chính (PK), khóa ngoại (FK) bảo đảm tính nhất quán ACID.\n"
                "• Đánh chỉ mục (Indexing): Tối ưu hóa truy vấn trên các trường slug, category_id, status."
            )
            set_text(sh, c_db, Pt(10), False, COLOR_TEXT_DARK)
        elif name == 'Text 7':
            set_text(sh, "Cấu trúc 15 bảng theo 4 nhóm thực thể chính:", Pt(11), True, COLOR_DARK_BLUE)
        elif name == 'Text 24':
            set_text(sh, "Sơ đồ quan hệ trong cơ sở dữ liệu", Pt(10), True, COLOR_TEXT_MUTED, PP_ALIGN.CENTER)
        elif name == 'Text 25':
            set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# ==========================================
# SLIDE 10: DEMO CỔNG THÔNG TIN
# ==========================================
s10 = prs.slides[9]
for sh in s10.shapes:
    name = sh.name
    if name == 'Text 0':
        set_text(sh, "5. KẾT QUẢ THỰC NGHIỆM", Pt(11), True, COLOR_GOLD)
    elif name == 'Text 1':
        set_text(sh, "Giao diện Cổng thông tin độc giả (User đọc báo & tra cứu)", Pt(18), True, COLOR_DARK_BLUE)
    elif name == 'Text 2':
        sh.left = 533400; sh.top = 1180000; sh.width = 2500000
        set_text(sh, "• Chuẩn nhận diện TDMU", Pt(10), True, COLOR_GOLD)
    elif name == 'Text 4':
        sh.left = 3200000; sh.top = 1180000; sh.width = 2800000
        set_text(sh, "• Kho Văn bản & Biểu mẫu", Pt(10), True, COLOR_GOLD)
    elif name == 'Text 6':
        sh.left = 6200000; sh.top = 1180000; sh.width = 2600000
        set_text(sh, "• Đọc bài viết chuẩn báo chí", Pt(10), True, COLOR_GOLD)
    elif name == 'Text 8':
        sh.left = 9000000; sh.top = 1180000; sh.width = 2600000
        set_text(sh, "• Tối ưu Responsive đa thiết bị", Pt(10), True, COLOR_GOLD)
    elif name in ['Text 3', 'Text 5', 'Text 7', 'Text 9']:
        hide_shape(sh)
    elif name == 'Text 11':
        set_text(sh, "Trang chủ Cổng thông tin", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 12':
        set_text(sh, "Bố cục trang nhã theo chuẩn nhận diện TDMU; menu chuyên mục rõ ràng, tích hợp banner sự kiện nổi bật và luồng tin phong trào mới nhất.", Pt(9.5), False, COLOR_TEXT_DARK)
    elif name == 'Text 14':
        set_text(sh, "Trang chi tiết bài viết & Tương tác", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 15':
        set_text(sh, "Định dạng bài báo chuẩn mực (Tiêu đề, Sapo, Toàn văn); hiển thị album ảnh sự kiện, tệp văn bản đính kèm và khu vực phản hồi độc giả.", Pt(9.5), False, COLOR_TEXT_DARK)
    elif name == 'Text 17':
        set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# ==========================================
# SLIDE 11: DEMO CMS & AI
# ==========================================
s11 = prs.slides[10]
for sh in s11.shapes:
    name = sh.name
    if name == 'Text 0':
        set_text(sh, "5. KẾT QUẢ THỰC NGHIỆM", Pt(11), True, COLOR_GOLD)
    elif name == 'Text 1':
        set_text(sh, "Phân hệ Quản trị CMS & Phòng biên tập tạo tin AI đa kênh", Pt(18), True, COLOR_DARK_BLUE)
    elif name == 'Text 2':
        sh.left = 533400; sh.top = 1180000; sh.width = 2500000
        set_text(sh, "• Bảng điều khiển Dashboard", Pt(10), True, COLOR_GOLD)
    elif name == 'Text 4':
        sh.left = 3200000; sh.top = 1180000; sh.width = 2800000
        set_text(sh, "• Biên tập viên quản lý xuất bản", Pt(10), True, COLOR_GOLD)
    elif name == 'Text 6':
        sh.left = 6200000; sh.top = 1180000; sh.width = 2600000
        set_text(sh, "• Cộng tác viên soạn bài AI", Pt(10), True, COLOR_GOLD)
    elif name == 'Text 8':
        sh.left = 9000000; sh.top = 1180000; sh.width = 2600000
        set_text(sh, "• Xuất bản đa kênh tự động", Pt(10), True, COLOR_GOLD)
    elif name in ['Text 3', 'Text 5', 'Text 7', 'Text 9']:
        hide_shape(sh)
    elif name == 'Text 11':
        set_text(sh, "Phòng Biên tập tích hợp Trợ lý AI", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 12':
        set_text(sh, "Soạn thảo bài viết trực quan; Trợ lý AI hỗ trợ tạo bài, gợi ý tiêu đề và xem trước định dạng bài đăng Fanpage Facebook, Zalo.", Pt(9.5), False, COLOR_TEXT_DARK)
    elif name == 'Text 14':
        set_text(sh, "Bảng Điều Hành (Dashboard)", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 15':
        set_text(sh, "Quản lý tập trung bài viết, lịch hẹn giờ đăng và theo dõi tổng thể số liệu tương tác đa kênh của hệ thống.", Pt(9.5), False, COLOR_TEXT_DARK)
    elif name == 'Text 17':
        set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

# ==========================================
# SLIDE 12: ĐÁNH GIÁ & HƯỚNG PHÁT TRIỂN
# ==========================================
s12 = prs.slides[11]
for sh in s12.shapes:
    name = sh.name
    if name == 'Text 0':
        set_text(sh, "6. ĐÁNH GIÁ & HƯỚNG PHÁT TRIỂN", Pt(11), True, COLOR_GOLD)
    elif name == 'Text 1':
        set_text(sh, "Kết quả đạt được và Hướng phát triển đề tài", Pt(18), True, COLOR_DARK_BLUE)
    elif name == 'Text 42':
        set_text(sh, "Nhóm 2 · Đồ án Cơ sở ngành 2026–2027", Pt(9), False, COLOR_TEXT_MUTED)

    elif name in ['Shape 11', 'Shape 12', 'Shape 13', 'Shape 14', 'Shape 19', 'Shape 21', 'Shape 22']:
        hide_shape(sh)

    elif name == 'Text 10':
        sh.left = 533400; sh.width = 5400000; sh.top = 1200000; sh.height = 1450000
    elif name == 'Text 16':
        sh.left = 733400; sh.width = 5000000; sh.top = 1350000
        set_text(sh, "TẠO BÀI VIẾT BẰNG AI & CHUẨN HÓA NỘI DUNG", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 17':
        sh.left = 733400; sh.width = 5000000; sh.top = 1650000; sh.height = 850000
        t1 = "Trợ lý AI tự động phân tích tư liệu sự kiện, hỗ trợ lập dàn ý và dựng bài viết hoàn chỉnh 5W1H; chuẩn hóa câu từ và đề xuất tiêu đề báo chí, rút ngắn đáng kể thời gian chuẩn bị tin bài cho cán bộ truyền thông."
        set_text(sh, t1, Pt(10), False, COLOR_TEXT_DARK)

    elif name == 'Text 18':
        sh.left = 6250000; sh.width = 5400000; sh.top = 1200000; sh.height = 1450000
    elif name == 'Text 24':
        sh.left = 6450000; sh.width = 5000000; sh.top = 1350000
        set_text(sh, "XUẤT BẢN ĐA KÊNH & LẬP LỊCH TỰ ĐỘNG", Pt(12), True, COLOR_DARK_BLUE)
    elif name == 'Text 25':
        sh.left = 6450000; sh.width = 5000000; sh.top = 1650000; sh.height = 850000
        t2 = "Tự động đăng bài và hình ảnh lên Fanpage Facebook đồng bộ cùng Website; cơ chế lập lịch ngầm chính xác, tin cậy; bảng điều khiển Dashboard theo dõi tổng hợp lượt xem và mức độ tương tác đa kênh."
        set_text(sh, t2, Pt(10), False, COLOR_TEXT_DARK)

    elif name == 'Text 35':
        sh.left = 533400; sh.width = 11125200; sh.top = 2850000; sh.height = 2400000
    elif name == 'Text 36':
        sh.left = 833400; sh.width = 10500000; sh.top = 3050000
        set_text(sh, "HƯỚNG PHÁT TRIỂN ĐỀ TÀI", Pt(13), True, COLOR_GOLD)
    elif name == 'Text 37':
        sh.left = 833400; sh.width = 10500000; sh.top = 3450000; sh.height = 1600000
        future_text = (
            "• Tinh chỉnh và tối ưu hóa bộ prompt chuyên sâu cho từng loại văn bản và hoạt động công đoàn đặc thù.\n\n"
            "• Hoàn tất kết nối Webhook chính thức với Zalo Official Account khi được cấp phép vận hành.\n\n"
            "• Nâng cấp các tiện ích khảo sát và lắng nghe tâm tư nguyện vọng của đoàn viên cơ sở trực tuyến.\n\n"
            "• Triển khai lấy ý kiến thử nghiệm thực tế từ các Tổ công đoàn để tối ưu trải nghiệm người dùng."
        )
        set_text(sh, future_text, Pt(10.5), False, COLOR_WHITE)

# ==========================================
# SLIDE 13: LỜI CẢM ƠN
# ==========================================
s13 = prs.slides[12]
for idx, sh in enumerate(s13.shapes):
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 39':
            sh.left = 1500000; sh.width = 9200000; sh.top = 2200000
            set_text(sh, "CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ LẮNG NGHE!", Pt(24), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif idx == 4 or 'Văn bản' in name or 'Text 2' in name:
            sh.left = 1500000; sh.width = 9200000; sh.top = 3600000; sh.height = 1500000
            thanks = (
                "Nhóm 2 xin trân trọng cảm ơn sự hướng dẫn tận tình của ThS. Võ Quốc Lương\n"
                "và những ý kiến đóng góp quý báu từ Quý Thầy Cô Viện Công nghệ số.\n\n"
                "Kính mời Quý Thầy Cô và các bạn nhận xét, đặt câu hỏi cho nhóm!"
            )
            set_text(sh, thanks, Pt(14), False, COLOR_TEXT_DARK, PP_ALIGN.CENTER)

# Save
prs.save(dst)
print("SUCCESS: 100% Perfect master alignment saved to thuyet trinh.pptx!")
