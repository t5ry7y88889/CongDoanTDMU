import pptx, sys
sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation(r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh_backup_before_enhancement.pptx')

with open(r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\orig_shapes_full.txt', 'w', encoding='utf-8') as f:
    for i, slide in enumerate(prs.slides):
        f.write(f'=== SLIDE {i+1} ===\n')
        for idx, sh in enumerate(slide.shapes):
            if sh.has_text_frame:
                tf = sh.text_frame
                full_txt = tf.text.strip().replace('\n', ' // ')
                if full_txt:
                    p0 = tf.paragraphs[0]
                    fz = p0.font.size.pt if p0.font.size else 'None'
                    b = p0.font.bold
                    color = None
                    try:
                        color = str(p0.font.color.rgb)
                    except:
                        pass
                    f.write(f'  [{idx}] {sh.name} (id={sh.shape_id}, pt={fz}, b={b}, c={color}) -> "{full_txt}"\n')
