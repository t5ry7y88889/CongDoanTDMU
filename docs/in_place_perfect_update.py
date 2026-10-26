import sys, os, shutil
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

sys.stdout.reconfigure(encoding='utf-8')

src = r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh_backup_before_enhancement.pptx'
dst = r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh.pptx'

# First, restore thuyet trinh.pptx directly from the clean original backup!
shutil.copy2(src, dst)
print("Restored original file cleanly!")

prs = Presentation(dst)

COLOR_DARK_BLUE = RGBColor(11, 44, 88)
COLOR_GOLD = RGBColor(218, 142, 17)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_TEXT_DARK = RGBColor(33, 37, 41)
COLOR_TEXT_MUTED = RGBColor(108, 117, 125)

def update_text(sh, text, font_size=None, bold=None, color=None, align=None):
    if not sh.has_text_frame:
        return
    tf = sh.text_frame
    tf.word_wrap = True
    
    # Save original font properties from first paragraph if available
    orig_font_name = "Calibri"
    orig_font_size = None
    orig_bold = None
    orig_color = None
    orig_align = None
    if len(tf.paragraphs) > 0:
        p0 = tf.paragraphs[0]
        if p0.font.name: orig_font_name = p0.font.name
        if p0.font.size: orig_font_size = p0.font.size
        orig_bold = p0.font.bold
        orig_align = p0.alignment
        try:
            if p0.font.color and p0.font.color.rgb:
                orig_color = p0.font.color.rgb
        except:
            pass

    tf.text = text
    
    # Apply styling
    for p in tf.paragraphs:
        p.font.name = orig_font_name
        p.font.size = font_size if font_size is not None else (orig_font_size if orig_font_size else Pt(10))
        p.font.bold = bold if bold is not None else orig_bold
        if color is not None:
            p.font.color.rgb = color
        elif orig_color is not None:
            p.font.color.rgb = orig_color
        if align is not None:
            p.alignment = align
        elif orig_align is not None:
            p.alignment = orig_align

# ==================== SLIDE 1 ====================
s1 = prs.slides[0]
for sh in s1.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 2':
            update_text(sh, "BÁO CÁO ĐỒ ÁN CƠ SỞ NGÀNH", Pt(13), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 4':
            update_text(sh, "GVHD: ThS. Võ Quốc Lương   |   Nhóm thực hiện: Nhóm 2", Pt(12), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif name == 'Text 5' and sh.shape_id == 2:
            members = (
                "1. Nguyễn Bình Dương (Trưởng nhóm) - MSSV: 2424802010319 - Lớp: D24CNTT05\n"
                "2. Trần Hồng Thanh - MSSV: 2424802010439 - Lớp: D24CNTT03\n"
                "3. Phạm Anh Tuấn - MSSV: 2324802010393 - Lớp: D23CNTT03"
            )
            update_text(sh, members, Pt(11), False, COLOR_TEXT_DARK, PP_ALIGN.CENTER)

# ==================== SLIDE 2 ====================
s2 = prs.slides[1]
for sh in s2.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Bối cảnh & Tính cấp thiết của đề tài")
        elif name == 'Text 4':
            update_text(sh, "Quy mô hơn 800 cán bộ, giảng viên và người lao động", Pt(10))
        elif name == 'Text 9':
            update_text(sh, "Kênh truyền thông hiện tại:", Pt(10), True)
        elif name == 'Text 10':
            update_text(sh, "Mục tin trên Website trường · Fanpage Facebook · Nhóm Zalo", Pt(9.5))
        elif name == 'Text 11':
            update_text(sh, "Đội ngũ truyền thông hoạt động kiêm nhiệm", Pt(10), True)
        elif name == 'Text 15':
            update_text(sh, "Áp lực kiêm nhiệm & Chậm trễ tin tức", Pt(11), True)
        elif name == 'Text 16':
            update_text(sh, "Cán bộ truyền thông là giảng viên kiêm nhiệm; khâu thu thập tư liệu và viết bài chuẩn phong cách báo chí đoàn thể tốn nhiều thời gian.", Pt(9.5))
        elif name == 'Text 20':
            update_text(sh, "Đăng tải thủ công, phân tán nhiều kênh", Pt(11), True)
        elif name == 'Text 21':
            update_text(sh, "Quy trình xuất bản rời rạc; cán bộ phải copy-paste văn bản và hình ảnh lặp lại thủ công trên từng nền tảng, dễ nhầm lẫn thông tin.", Pt(9.5))
        elif name == 'Text 25':
            update_text(sh, "Thiếu công cụ đo lường tập trung", Pt(11), True)
        elif name == 'Text 26':
            update_text(sh, "Ban Chấp hành chưa có bảng số liệu đo lường định lượng tương tác đa kênh để đánh giá và định hướng hiệu quả thông tin đoàn thể.", Pt(9.5))

# ==================== SLIDE 3 ====================
s3 = prs.slides[2]
for sh in s3.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Mục tiêu nghiên cứu & Định hướng giải quyết")
        elif name == 'Text 9':
            update_text(sh, "Mục tiêu trọng tâm của đề tài", Pt(12), True)
        elif name == 'Text 10':
            t10 = (
                "• Xây dựng hệ thống website truyền thông chuyên biệt của Công đoàn Trường ĐH Thủ Dầu Một.\n"
                "• Tích hợp Trợ lý AI hỗ trợ tạo bài viết từ tư liệu sự kiện thô và tự động hóa xuất bản đa kênh.\n"
                "• Giải phóng sức lao động thủ công, chuẩn hóa thông tin và nâng cao hiệu quả tuyên truyền đoàn thể."
            )
            update_text(sh, t10, Pt(10), False)
        elif name == 'Text 18':
            update_text(sh, "Hỗ trợ tạo nội dung bằng AI", Pt(12), True)
        elif name == 'Text 19':
            t19 = (
                "• Phân tích kế hoạch, thông báo sự kiện để dựng bài 5W1H.\n"
                "• Tạo tóm tắt Sapo thu hút, gợi ý tiêu đề báo chí chuẩn SEO.\n"
                "• Chuẩn hóa văn phong, rà soát chính tả tiếng Việt.\n"
                "• Tối ưu định dạng ngắn kèm hashtag cho Facebook và Zalo."
            )
            update_text(sh, t19, Pt(9.5), False)
        elif name == 'Text 27':
            update_text(sh, "Xuất bản đa kênh & Đo lường", Pt(12), True)
        elif name == 'Text 28':
            t28 = (
                "• Tự động đăng đồng thời lên Website và Fanpage Facebook.\n"
                "• Cơ chế lập lịch hẹn giờ xuất bản ngầm chính xác, tin cậy.\n"
                "• Phân định 4 vai trò rõ rệt: Admin, Editor, Contributor, User.\n"
                "• Dashboard theo dõi định lượng lượt xem và độ lan tỏa tin bài."
            )
            update_text(sh, t28, Pt(9.5), False)

# ==================== SLIDE 4 ====================
s4 = prs.slides[3]
for sh in s4.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Cổng thông tin độc giả — Đọc báo, tra cứu & tương tác")
        elif name == 'Text 5':
            update_text(sh, "Tin tức phong trào & Phân loại đa tầng", Pt(11), True)
        elif name == 'Text 6':
            update_text(sh, "Phân loại theo 8 chuyên mục công đoàn; tìm kiếm toàn văn theo từ khóa, tác giả, thời gian; phân trang mượt mà và hiển thị tin nổi bật trên Slider trang chủ.", Pt(9.5))
        elif name == 'Text 10':
            update_text(sh, "Bài viết đa phương tiện chuẩn báo chí", Pt(11), True)
        elif name == 'Text 11':
            update_text(sh, "Bố cục báo chí chuyên nghiệp (Tiêu đề, Sapo cô đọng, Thân bài rich-text); hiển thị album ảnh sự kiện độ phân giải cao và tệp văn bản đính kèm xem trực tiếp.", Pt(9.5))
        elif name == 'Text 15':
            update_text(sh, "Tương tác độc giả & Chia sẻ liên kết", Pt(11), True)
        elif name == 'Text 16':
            update_text(sh, "Đoàn viên bày tỏ cảm xúc (Like), bình luận trao đổi văn minh; chia sẻ một chạm liên kết bài viết sang Facebook, Zalo cá nhân; gửi thư góp ý trực tuyến.", Pt(9.5))
        elif name == 'Text 20':
            update_text(sh, "Kho văn bản pháp quy & Biểu mẫu số hóa", Pt(11), True)
        elif name == 'Text 21':
            update_text(sh, "Tra cứu hệ thống nghị quyết, quy chế công đoàn minh bạch; tải về biểu mẫu xin trợ cấp khó khăn, mẫu đề xuất khen thưởng định dạng Word/PDF chuẩn hóa.", Pt(9.5))

# ==================== SLIDE 5 ====================
s5 = prs.slides[4]
for sh in s5.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Phân hệ Quản trị CMS & Điều phối xuất bản đa kênh")
        elif name == 'Text 5':
            update_text(sh, "Phân quyền 4 tác nhân", Pt(10.5), True)
        elif name == 'Text 6':
            update_text(sh, "Admin quản trị tài khoản; Editor quản lý tin bài & xuất bản; Contributor viết bài gửi lên; User đọc báo và tra cứu.", Pt(9))
        elif name == 'Text 10':
            update_text(sh, "Soạn thảo bài viết trực quan", Pt(10.5), True)
        elif name == 'Text 11':
            update_text(sh, "Trình soạn thảo rich-text hiện đại, hỗ trợ chèn bảng biểu, định dạng báo chí, tải album ảnh sự kiện và tệp đính kèm.", Pt(9))
        elif name == 'Text 15':
            update_text(sh, "Quản lý trạng thái bài viết", Pt(10.5), True)
        elif name == 'Text 16':
            update_text(sh, "Theo dõi vòng đời bài viết: Bản nháp → Đã gửi → Xuất bản ngay / Hẹn giờ phát hành; xem trước giao diện hiển thị.", Pt(9))
        elif name == 'Text 20':
            update_text(sh, "Lập lịch hẹn giờ tự động", Pt(10.5), True)
        elif name == 'Text 21':
            update_text(sh, "Thiết lập thời gian đăng bài linh hoạt; tiến trình ngầm (Node-cron) tự động kích hoạt xuất bản đúng giờ mà không cần trực máy.", Pt(9))
        elif name == 'Text 25':
            update_text(sh, "Đồng bộ Fanpage Facebook", Pt(10.5), True)
        elif name == 'Text 26':
            update_text(sh, "Tích hợp Meta Graph API v19 tự động đăng tải bài viết tóm tắt, hình ảnh đại diện và liên kết đọc báo lên Fanpage Công đoàn.", Pt(9))
        elif name == 'Text 30':
            update_text(sh, "Bảng điều khiển (Dashboard)", Pt(10.5), True)
        elif name == 'Text 31':
            update_text(sh, "Theo dõi trực quan tổng số tin bài, trạng thái xuất bản và biểu đồ định lượng mức độ tương tác độc giả đa kênh.", Pt(9))
        elif name == 'Text 35':
            update_text(sh, "Quản lý Văn bản & Biểu mẫu", Pt(10.5), True)
        elif name == 'Text 36':
            update_text(sh, "Cập nhật số hiệu văn bản, trích yếu, tệp scan PDF và phân loại các biểu mẫu đơn điện tử theo danh mục rõ ràng.", Pt(9))
        elif name == 'Text 40':
            update_text(sh, "Quản trị hệ thống & AI Token", Pt(10.5), True)
        elif name == 'Text 41':
            update_text(sh, "Quản lý tài khoản người dùng, phân quyền RBAC, quản lý an toàn khóa API (Gemini, Facebook) và định mức Token AI.", Pt(9))

# ==================== SLIDE 6 ====================
s6 = prs.slides[5]
for sh in s6.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Trợ lý AI hỗ trợ sáng tạo nội dung đa kênh")
        elif name == 'Text 2':
            update_text(sh, "Ứng dụng Google Gemini AI đọc hiểu tư liệu thô, tự động tạo bài viết và tối ưu định dạng phát hành theo từng kênh.", Pt(10))
        elif name == 'Text 6':
            update_text(sh, "Chắp bút từ tư liệu sự kiện", Pt(11), True)
        elif name == 'Text 7':
            update_text(sh, "Đọc hiểu kế hoạch, thông báo, ảnh chụp sự kiện (DOCX, PDF, ảnh scan) để dựng bản thảo hoàn chỉnh theo cấu trúc 5W1H báo chí đoàn thể.", Pt(9.5))
        elif name == 'Text 11':
            update_text(sh, "Tóm tắt Sapo & Đa kênh", Pt(11), True)
        elif name == 'Text 12':
            update_text(sh, "Tự động tạo đoạn mở đầu (Sapo) cô đọng cho Website; rút gọn thành bản tin ngắn kèm hashtag phù hợp đăng Facebook và Zalo.", Pt(9.5))
        elif name == 'Text 25':
            update_text(sh, "03", Pt(11), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 26':
            update_text(sh, "Gợi ý tiêu đề bài viết chuẩn SEO", Pt(11), True)
        elif name == 'Text 27':
            update_text(sh, "Phân tích nội dung cốt lõi và đề xuất 3–5 phương án tiêu đề trang trọng, đúng văn phong công đoàn và thu hút độc giả.", Pt(9.5))
        elif name == 'Text 30':
            update_text(sh, "04", Pt(11), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 31':
            update_text(sh, "Chuẩn hóa văn phong & Ngữ pháp", Pt(11), True)
        elif name == 'Text 32':
            update_text(sh, "Rà soát chính tả tiếng Việt, tinh chỉnh câu từ mạch lạc, thay thế từ ngữ khẩu ngữ bằng thuật ngữ chuẩn mực văn bản hành chính đoàn thể.", Pt(9.5))
        elif name == 'Text 34':
            update_text(sh, "Hạ tầng kỹ thuật: Tích hợp Google Gemini API qua giao thức an toàn, xử lý dữ liệu thời gian thực và quản lý hạn mức Token.", Pt(10), True, COLOR_WHITE, PP_ALIGN.CENTER)

# ==================== SLIDE 7 ====================
s7 = prs.slides[6]
for sh in s7.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Kiến trúc kỹ thuật — 4 thành phần hệ thống")
        elif name == 'Text 11':
            update_text(sh, "React 18 SPA · Tailwind CSS", Pt(10), True)
        elif name == 'Text 14':
            update_text(sh, "Cổng thông tin và CMS cơ chế tải nhanh Single Page Application, chuẩn nhận diện thương hiệu TDMU, tối ưu hiển thị Responsive trên PC và di động.", Pt(9.5))
        elif name == 'Text 22':
            update_text(sh, "Node.js 18+ · Express", Pt(10), True)
        elif name == 'Text 24':
            update_text(sh, "RESTful API", Pt(9.5), True)
        elif name == 'Text 28':
            update_text(sh, "Meta Graph API v19", Pt(9.5), True)
        elif name == 'Text 29':
            update_text(sh, "Xử lý nghiệp vụ tập trung, xác thực bảo mật JWT, mã hóa Bcrypt và điều phối tiến trình ngầm (Node-cron) tự động lập lịch xuất bản tin bài đúng giờ.", Pt(9.5))
        elif name == 'Text 37':
            update_text(sh, "Microsoft SQL Server 2022", Pt(10), True)
        elif name == 'Text 42':
            update_text(sh, "Mô hình dữ liệu quan hệ chuẩn hóa 3NF, bảo đảm toàn vẹn giao dịch ACID, tối ưu hóa truy vấn thông qua hệ thống chỉ mục (Indexing) và quan hệ chặt chẽ.", Pt(9.5))
        elif name == 'Text 50':
            update_text(sh, "Google Gemini API", Pt(10), True)
        elif name == 'Text 53':
            update_text(sh, "Meta Graph API v19", Pt(10), True)
        elif name == 'Text 54':
            update_text(sh, "Gemini xử lý ngôn ngữ tự nhiên hỗ trợ tạo bài viết; Meta Graph API cung cấp kết nối an toàn tự động đăng bài lên Fanpage ngay khi có lệnh xuất bản.", Pt(9.5))

# ==================== SLIDE 8 ====================
s8 = prs.slides[7]
for sh in s8.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Biểu đồ Use Case tổng quát và 4 nhóm tác nhân")
        elif name == 'Text 3':
            update_text(sh, "Quản trị viên (Admin)", Pt(9.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 4':
            update_text(sh, "Biên tập viên (Editor)", Pt(9.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 5':
            update_text(sh, "Cộng tác viên (Contributor)", Pt(9.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 6':
            update_text(sh, "Đoàn viên / Độc giả (User)", Pt(9.5), True, COLOR_WHITE, PP_ALIGN.CENTER)
        elif name == 'Text 7':
            update_text(sh, "5 PHÂN HỆ NGHIỆP VỤ (26 USE CASE)", Pt(11), True)
        elif name == 'Text 11':
            update_text(sh, "Quản lý bài viết & nội dung — Soạn thảo, đính kèm media, lưu nháp, phân quyền tác giả.", Pt(9.5))
        elif name == 'Text 14':
            update_text(sh, "Trợ lý AI hỗ trợ sáng tạo — Dựng bài 5W1H, gợi ý tiêu đề, tóm tắt Sapo.", Pt(9.5))
        elif name == 'Text 17':
            update_text(sh, "Xuất bản đa kênh & Lập lịch — Xem trước, hẹn giờ đăng bài, đồng bộ Facebook.", Pt(9.5))
        elif name == 'Text 23':
            update_text(sh, "Cổng thông tin độc giả — Đọc tin tức, tra cứu văn bản pháp quy, gửi góp ý.", Pt(9.5))
        elif name == 'Text 25':
            update_text(sh, "Quản trị hệ thống & Người dùng — Quản lý tài khoản, phân quyền RBAC, Token AI.", Pt(9.5))
        elif name == 'Text 27':
            note8 = "Phân định 4 vai trò rõ rệt: Admin quản trị hệ thống; Editor duyệt bài & xuất bản; Contributor soạn & gửi bài; User đọc & tương tác. Toàn bộ 26 use case được đặc tả chuẩn UML."
            update_text(sh, note8, Pt(9.5))

# ==================== SLIDE 9 ====================
s9 = prs.slides[8]
for sh in s9.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Thiết kế cơ sở dữ liệu quan hệ (15 bảng chuẩn 3NF)")
        elif name == 'Text 3':
            update_text(sh, "Đặc tả kỹ thuật CSDL", Pt(11), True)
        elif name == 'Text 4':
            c_db = (
                "• Hệ quản trị: Microsoft SQL Server 2022 kết hợp Connection Pool tối ưu hiệu năng kết nối.\n"
                "• Chuẩn hóa dữ liệu: Thiết kế chuẩn bậc 3 (3NF), loại bỏ dư thừa và hạn chế dị thường dữ liệu.\n"
                "• Ràng buộc toàn vẹn: Khóa chính (PK), khóa ngoại (FK) bảo đảm tính nhất quán giao dịch ACID.\n"
                "• Đánh chỉ mục (Indexing): Tối ưu hóa truy vấn trên các trường slug, category_id, status."
            )
            update_text(sh, c_db, Pt(9.5))
        elif name == 'Text 6':
            update_text(sh, "15 bảng theo 4 nhóm thực thể chính:", Pt(10.5), True)

# ==================== SLIDE 10 ====================
s10 = prs.slides[9]
for sh in s10.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Giao diện Cổng thông tin độc giả (User đọc báo & tra cứu)")
        elif name == 'Text 2':
            update_text(sh, "Chuẩn nhận diện thương hiệu TDMU", Pt(9.5), True)
        elif name == 'Text 4':
            update_text(sh, "Bảng tin & Slider sự kiện phong trào", Pt(9.5), True)
        elif name == 'Text 6':
            update_text(sh, "Đọc bài viết đa phương tiện & Tương tác", Pt(9.5), True)
        elif name == 'Text 8':
            update_text(sh, "Tối ưu hiển thị Responsive đa thiết bị", Pt(9.5), True)
        elif name == 'Text 12':
            update_text(sh, "Header nhận diện thương hiệu, menu chuyên mục rõ ràng, tích hợp banner sự kiện nổi bật và luồng tin phong trào mới nhất.", Pt(9.5))
        elif name == 'Text 15':
            update_text(sh, "Định dạng bài báo chuẩn mực (Tiêu đề, Sapo, Toàn văn); hiển thị album ảnh sự kiện, tệp văn bản đính kèm và khu vực phản hồi độc giả.", Pt(9.5))

# ==================== SLIDE 11 ====================
s11 = prs.slides[10]
for sh in s11.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Phân hệ Quản trị CMS & Phòng biên tập tạo tin AI đa kênh")
        elif name == 'Text 2':
            update_text(sh, "Bảng điều khiển (Dashboard) theo dõi tin bài", Pt(9.5), True)
        elif name == 'Text 4':
            update_text(sh, "Biên tập viên quản lý xuất bản", Pt(9.5), True)
        elif name == 'Text 6':
            update_text(sh, "Cộng tác viên soạn thảo & dùng Trợ lý AI", Pt(9.5), True)
        elif name == 'Text 8':
            update_text(sh, "Xuất bản đa kênh tự động (Web, Facebook)", Pt(9.5), True)
        elif name == 'Text 11':
            update_text(sh, "Phòng Biên tập tích hợp Trợ lý AI", Pt(11), True)
        elif name == 'Text 12':
            update_text(sh, "Soạn thảo bài viết trực quan; Trợ lý AI hỗ trợ tạo bài, gợi ý tiêu đề và xem trước định dạng bài đăng Fanpage Facebook, Zalo.", Pt(9.5))
        elif name == 'Text 14':
            update_text(sh, "Bảng Điều Hành (Dashboard)", Pt(11), True)
        elif name == 'Text 15':
            update_text(sh, "Quản lý tập trung bài viết, lịch hẹn giờ đăng và theo dõi tổng thể số liệu tương tác đa kênh của hệ thống.", Pt(9.5))

# ==================== SLIDE 12 ====================
s12 = prs.slides[11]
for sh in s12.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            update_text(sh, "Kết quả đạt được và Hướng phát triển đề tài")
        elif name == 'Text 16':
            update_text(sh, "TẠO BÀI VIẾT BẰNG AI", Pt(11), True, COLOR_GOLD, PP_ALIGN.CENTER)
        elif name == 'Text 17':
            update_text(sh, "Trợ lý AI hỗ trợ dựng bài 5W1H từ tư liệu thô, gợi ý tiêu đề báo chí và chuẩn hóa văn phong công đoàn.", Pt(8.5), False, COLOR_TEXT_DARK, PP_ALIGN.CENTER)
        elif name == 'Text 24':
            update_text(sh, "XUẤT BẢN ĐA KÊNH & LẬP LỊCH", Pt(11), True, COLOR_GOLD, PP_ALIGN.CENTER)
        elif name == 'Text 25':
            update_text(sh, "Tự động đăng tải đồng thời lên Website và Facebook Fanpage; lập lịch xuất bản ngầm chính xác theo giờ hẹn.", Pt(8.5), False, COLOR_TEXT_DARK, PP_ALIGN.CENTER)
        elif name == 'Text 36':
            update_text(sh, "HƯỚNG PHÁT TRIỂN ĐỀ TÀI", Pt(12), True, COLOR_GOLD)
        elif name == 'Text 37':
            future = (
                "• Tinh chỉnh và tối ưu hóa bộ prompt chuyên sâu cho từng loại văn bản và hoạt động công đoàn đặc thù.\n"
                "• Hoàn tất kết nối Webhook chính thức với Zalo Official Account khi được cấp phép vận hành.\n"
                "• Nâng cấp các tiện ích khảo sát và lắng nghe tâm tư nguyện vọng của đoàn viên cơ sở trực tuyến.\n"
                "• Triển khai lấy ý kiến thử nghiệm thực tế từ các Tổ công đoàn để tối ưu hóa trải nghiệm người dùng."
            )
            update_text(sh, future, Pt(9.5), False, COLOR_WHITE)

# ==================== SLIDE 13 ====================
s13 = prs.slides[12]
for sh in s13.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 39':
            update_text(sh, "CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ LẮNG NGHE!", Pt(22), True, COLOR_DARK_BLUE, PP_ALIGN.CENTER)
        elif 'Văn bản' in name or 'Hộp' in name:
            thanks = (
                "Nhóm 2 xin trân trọng cảm ơn sự hướng dẫn tận tình của ThS. Võ Quốc Lương\n"
                "và những ý kiến đóng góp quý báu từ Quý Thầy Cô Viện Công nghệ số.\n\n"
                "Kính mời Quý Thầy Cô và các bạn nhận xét, đặt câu hỏi cho nhóm!"
            )
            update_text(sh, thanks, Pt(13), False, COLOR_TEXT_DARK, PP_ALIGN.CENTER)

prs.save(dst)
print("SUCCESS: 100% In-Place text update done without touching any shape positions or layout!")
