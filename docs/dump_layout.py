import sys, os
from pptx import Presentation

prs = Presentation(r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh_khoi_phuc_goc.pptx')
with open(r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\orig_layout_dump.txt', 'w', encoding='utf-8') as f:
    for i, s in enumerate(prs.slides):
        f.write(f'=== SLIDE {i+1} ===\n')
        for j, sh in enumerate(s.shapes):
            txt = ''
            font_size = ''
            if sh.has_text_frame:
                txt = sh.text_frame.text.strip().replace('\n', ' // ')[:60]
                if len(sh.text_frame.paragraphs) > 0 and sh.text_frame.paragraphs[0].font.size:
                    font_size = f'{sh.text_frame.paragraphs[0].font.size.pt}pt'
            f.write(f'  [{j}] {sh.name} type={sh.shape_type} pos=({sh.left},{sh.top},{sh.width},{sh.height}) font={font_size} txt="{txt}"\n')
print('Dumped successfully!')
