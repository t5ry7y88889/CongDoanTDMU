import sys, os
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding='utf-8')

src = r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh_khoi_phuc_goc.pptx'
dst = r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh.pptx'

prs = Presentation(src)

def set_shape_text(sh, text, font_size=None, bold=None, color=None):
    if not sh.has_text_frame:
        return
    sh.text_frame.text = text
    sh.text_frame.word_wrap = True
    for p in sh.text_frame.paragraphs:
        if font_size:
            p.font.size = font_size
        if bold is not None:
            p.font.bold = bold
        if color:
            p.font.color.rgb = color

# ==================== SLIDE 1 ====================
s1 = prs.slides[0]
for sh in s1.shapes:
    if sh.has_text_frame:
        name = sh.name
        t = sh.text_frame.text.strip()
        if name == 'Text 2':
            set_shape_text(sh, 'BÁO CÁO ĐỒ ÁN CƠ SỞ NGÀNH', Pt(16), True)
        elif name == 'Text 4':
            set_shape_text(sh, 'GVHD: ThS. Võ Quốc Lương  ·  Nhóm thực hiện: Nhóm 2', Pt(12), True)
        elif name == 'Text 5' and sh.top > 4000000:
            set_shape_text(sh, '1. Nguyễn Bình Dương (Trưởng nhóm) - MSSV: 2424802010319\n2. Trần Hồng Thanh - MSSV: 2424802010439\n3. Phạm Anh Tuấn - MSSV: 2324802010393', Pt(10.5), False)

# ==================== SLIDE 2 ====================
# Slide 2 in original has 3 cards:
# Card 1: Text 14 ("01"), Text 15 (Title), Text 16 (Desc)
# Card 2: Text 19 ("02"), Text 20 (Title), Text 21 (Desc)
# Card 3: Text 24 ("03"), Text 25 (Title), Text 26 (Desc)
s2 = prs.slides[1]
for sh in s2.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            set_shape_text(sh, 'Bối cảnh & Tính cấp thiết của đề tài', Pt(18), True)
        elif name == 'Text 4':
            set_shape_text(sh, 'Quy mô hơn 800 cán bộ, giảng viên và người lao động', Pt(10.5))
        elif name == 'Text 9':
            set_shape_text(sh, 'Kênh truyền thông hiện tại:', Pt(10), True)
        elif name == 'Text 10':
            set_shape_text(sh, 'Mục tin trên Website trường · Fanpage Facebook · Nhóm Zalo', Pt(10))
        elif name == 'Text 11':
            set_shape_text(sh, 'Đội ngũ truyền thông hoạt động kiêm nhiệm', Pt(10))
        # Card 1
        elif name == 'Text 14':
            set_shape_text(sh, '01', Pt(14), True)
        elif name == 'Text 15':
            set_shape_text(sh, 'Áp lực kiêm nhiệm & Chậm trễ tin tức', Pt(11), True)
        elif name == 'Text 16':
            set_shape_text(sh, 'Cán bộ làm truyền thông là giảng viên kiêm nhiệm chuyên môn; khâu chuẩn bị tư liệu, viết bài chuẩn phong cách đoàn thể và tìm ảnh tốn nhiều thời gian.', Pt(9.5))
        # Card 2
        elif name == 'Text 19':
            set_shape_text(sh, '02', Pt(14), True)
        elif name == 'Text 20':
            set_shape_text(sh, 'Đăng tải thủ công, lặp lại nhiều kênh', Pt(11), True)
        elif name == 'Text 21':
            set_shape_text(sh, 'Quy trình xuất bản rời rạc; cán bộ phải sao chép văn bản, định dạng lại và tải ảnh lặp lại thủ công trên từng nền tảng (Website, Fanpage, Zalo).', Pt(9.5))
        # Card 3
        elif name == 'Text 24':
            set_shape_text(sh, '03', Pt(14), True)
        elif name == 'Text 25':
            set_shape_text(sh, 'Thiếu công cụ đo lường tập trung', Pt(11), True)
        elif name == 'Text 26':
            set_shape_text(sh, 'Ban Chấp hành chưa có công cụ số liệu đo lường định lượng lượt xem, mức độ tương tác đa kênh để đánh giá và điều chỉnh hiệu quả tuyên truyền.', Pt(9.5))

# ==================== SLIDE 3 ====================
# Box 1: Text 8 ("01"), Text 9 (Title), Text 10 (Desc)
# Box 2: Text 17 ("02"), Text 18 (Title), Text 19 (Desc)
# Box 3: Text 26 ("03"), Text 27 (Title), Text 28 (Desc)
s3 = prs.slides[2]
for sh in s3.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            set_shape_text(sh, 'Mục tiêu nghiên cứu & Định hướng giải quyết', Pt(18), True)
        # Box 1: Mục tiêu trọng tâm
        elif name == 'Text 9':
            set_shape_text(sh, 'Mục tiêu trọng tâm của đề tài', Pt(12.5), True)
        elif name == 'Text 10':
            set_shape_text(sh, 'Xây dựng hệ thống quản lý truyền thông tích hợp AI hỗ trợ tạo nội dung và tự động hóa xuất bản đa kênh (Website Công đoàn, Fanpage Facebook, Zalo OA), giải phóng sức lao động thủ công và nâng cao hiệu quả thông tin đoàn thể.', Pt(10.5))
        # Box 2: Tạo nội dung bằng AI
        elif name == 'Text 18':
            set_shape_text(sh, 'Hỗ trợ tạo nội dung bằng AI', Pt(11.5), True)
        elif name == 'Text 19':
            set_shape_text(sh, '• Tự động phân tích kế hoạch, tư liệu sự kiện để lập dàn ý và dựng bài 5W1H.\n• Tự động tạo tóm tắt Sapo, gợi ý tiêu đề thu hút độc giả.\n• Rà soát chính tả tiếng Việt và chuẩn hóa văn phong báo chí đoàn thể.\n• Chuyển đổi định dạng bài đăng ngắn kèm hashtag cho Facebook và Zalo.', Pt(9.5))
        # Box 3: Xuất bản đa kênh
        elif name == 'Text 27':
            set_shape_text(sh, 'Xuất bản đa kênh tự động & Đo lường', Pt(11.5), True)
        elif name == 'Text 28':
            set_shape_text(sh, '• Tự động đăng tải nội dung và hình ảnh đồng thời lên Website, Fanpage, Zalo.\n• Cơ chế lập lịch hẹn giờ xuất bản ngầm (Background Scheduler) chính xác.\n• Phân định 4 vai trò rõ rệt: Admin, Editor, Contributor, User.\n• Dashboard theo dõi định lượng lượt xem và tương tác đa kênh tập trung.', Pt(9.5))

# ==================== SLIDE 4 ====================
# 4 cards:
# Card 1: Text 4 ("01"), Text 5 (Title), Text 6 (Desc)
# Card 2: Text 9 ("02"), Text 10 (Title), Text 11 (Desc)
# Card 3: Text 14 ("03"), Text 15 (Title), Text 16 (Desc)
# Card 4: Text 19 ("04"), Text 20 (Title), Text 21 (Desc)
s4 = prs.slides[3]
for sh in s4.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            set_shape_text(sh, 'Cổng thông tin độc giả — Đọc báo, tra cứu & tương tác', Pt(18), True)
        elif name == 'Text 5':
            set_shape_text(sh, 'Tin tức phong trào & Tìm kiếm', Pt(11), True)
        elif name == 'Text 6':
            set_shape_text(sh, 'Phân loại tin theo chuyên mục hoạt động; tìm kiếm toàn văn theo từ khóa, tác giả, thời gian; danh sách hiển thị phân trang tối ưu.', Pt(9.5))
        elif name == 'Text 10':
            set_shape_text(sh, 'Đọc bài viết đa phương tiện', Pt(11), True)
        elif name == 'Text 11':
            set_shape_text(sh, 'Bố cục bài báo chuẩn mực (Tiêu đề – Sapo – Toàn văn); hiển thị album ảnh sự kiện chất lượng cao và tệp văn bản đính kèm.', Pt(9.5))
        elif name == 'Text 15':
            set_shape_text(sh, 'Tương tác & Chia sẻ đa nền tảng', Pt(11), True)
        elif name == 'Text 16':
            set_shape_text(sh, 'Đoàn viên bày tỏ cảm xúc (Like), bình luận trao đổi; chia sẻ bài viết nhanh chóng sang Facebook, Zalo cá nhân; gửi hòm thư góp ý.', Pt(9.5))
        elif name == 'Text 20':
            set_shape_text(sh, 'Kho văn bản & Biểu mẫu điện tử', Pt(11), True)
        elif name == 'Text 21':
            set_shape_text(sh, 'Tra cứu nghị quyết, quy định, luật công đoàn kèm tệp PDF scan; cung cấp mẫu đơn điện tử cho phép tải về định dạng Word/PDF.', Pt(9.5))

# ==================== SLIDE 5 ====================
# 8 cards:
# Card 1: Text 5 (Title), Text 6 (Desc)
# Card 2: Text 10 (Title), Text 11 (Desc)
# Card 3: Text 15 (Title), Text 16 (Desc)
# Card 4: Text 20 (Title), Text 21 (Desc)
# Card 5: Text 25 (Title), Text 26 (Desc)
# Card 6: Text 30 (Title), Text 31 (Desc)
# Card 7: Text 35 (Title), Text 36 (Desc)
# Card 8: Text 40 (Title), Text 41 (Desc)
s5 = prs.slides[4]
for sh in s5.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            set_shape_text(sh, 'Phân hệ Quản trị CMS & Điều phối xuất bản đa kênh', Pt(18), True)
        elif name == 'Text 5':
            set_shape_text(sh, 'Phân quyền 4 tác nhân rõ rệt', Pt(10.5), True)
        elif name == 'Text 6':
            set_shape_text(sh, 'Admin quản lý tài khoản; Editor quản lý tin bài & xuất bản; Contributor viết bài gửi lên; User đọc báo và tương tác.', Pt(9))
        elif name == 'Text 10':
            set_shape_text(sh, 'Soạn thảo bài viết trực quan', Pt(10.5), True)
        elif name == 'Text 11':
            set_shape_text(sh, 'Trình soạn thảo hiện đại, hỗ trợ định dạng chuẩn báo chí, bảng biểu và chèn tệp media.', Pt(9))
        elif name == 'Text 15':
            set_shape_text(sh, 'Quản lý trạng thái bài viết', Pt(10.5), True)
        elif name == 'Text 16':
            set_shape_text(sh, 'Theo dõi trạng thái: Bản nháp → Đã gửi → Xuất bản ngay / Hẹn giờ phát hành; xem trước giao diện hiển thị.', Pt(9))
        elif name == 'Text 20':
            set_shape_text(sh, 'Lập lịch xuất bản tự động', Pt(10.5), True)
        elif name == 'Text 21':
            set_shape_text(sh, 'Thiết lập thời gian phát hành bài viết theo khung giờ ấn định; tiến trình ngầm tự động kích hoạt xuất bản đúng giờ.', Pt(9))
        elif name == 'Text 25':
            set_shape_text(sh, 'Đồng bộ xuất bản Facebook', Pt(10.5), True)
        elif name == 'Text 26':
            set_shape_text(sh, 'Tích hợp Meta Graph API v19 tự động đăng bài viết tóm tắt, hình ảnh và liên kết bài đọc lên Fanpage Công đoàn.', Pt(9))
        elif name == 'Text 30':
            set_shape_text(sh, 'Bảng điều khiển (Dashboard)', Pt(10.5), True)
        elif name == 'Text 31':
            set_shape_text(sh, 'Theo dõi tổng quan số lượng bài viết, trạng thái phát hành và biểu đồ mức độ tương tác đa kênh.', Pt(9))
        elif name == 'Text 35':
            set_shape_text(sh, 'Quản lý Văn bản & Biểu mẫu', Pt(10.5), True)
        elif name == 'Text 36':
            set_shape_text(sh, 'Cập nhật số hiệu văn bản, trích yếu, tệp scan PDF; phân loại biểu mẫu đơn từ điện tử theo danh mục.', Pt(9))
        elif name == 'Text 40':
            set_shape_text(sh, 'Quản trị hệ thống & Token AI', Pt(10.5), True)
        elif name == 'Text 41':
            set_shape_text(sh, 'Quản lý tài khoản người dùng, phân quyền RBAC, cấu hình khóa API (Gemini, Facebook) và định mức Token AI.', Pt(9))

# ==================== SLIDE 6 ====================
# Card 1: Text 5 ("01"), Text 6 (Title), Text 7 (Desc)
# Card 2: Text 10 ("02"), Text 11 (Title), Text 12 (Desc)
# Card 3: Text 25 ("03"), Text 26 (Title), Text 27 (Desc)
# Card 4: Text 30 ("04"), Text 31 (Title), Text 32 (Desc)
s6 = prs.slides[5]
for sh in s6.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            set_shape_text(sh, 'Trợ lý AI hỗ trợ sáng tạo nội dung đa kênh', Pt(18), True)
        elif name == 'Text 2':
            set_shape_text(sh, 'Ứng dụng Google Gemini AI đọc hiểu tư liệu thô, tự động tạo bài viết và tối ưu định dạng phát hành cho từng kênh (Website, Fanpage, Zalo).', Pt(10))
        # Card 1
        elif name == 'Text 6':
            set_shape_text(sh, 'Chấp bút bài viết từ tư liệu thô', Pt(11), True)
        elif name == 'Text 7':
            set_shape_text(sh, 'Đọc hiểu kế hoạch, thông báo, ảnh sự kiện để tự động dựng bài viết hoàn chỉnh theo cấu trúc chuẩn 5W1H báo chí.', Pt(9.5))
        # Card 2
        elif name == 'Text 11':
            set_shape_text(sh, 'Tóm tắt Sapo & Định dạng đa kênh', Pt(11), True)
        elif name == 'Text 12':
            set_shape_text(sh, 'Tự động tạo đoạn mở đầu Sapo cho Website, đồng thời rút gọn thành bài đăng ngắn kèm hashtag cho Facebook và Zalo.', Pt(9.5))
        # Card 3 (sửa số 04 -> 03)
        elif name == 'Text 25':
            set_shape_text(sh, '03', Pt(14), True)
        elif name == 'Text 26':
            set_shape_text(sh, 'Gợi ý tiêu đề bài viết', Pt(11), True)
        elif name == 'Text 27':
            set_shape_text(sh, 'Phân tích nội dung cốt lõi và đề xuất 3–5 phương án tiêu đề trang trọng, đúng văn phong đoàn thể và thu hút độc giả.', Pt(9.5))
        # Card 4 (sửa số 05 -> 04)
        elif name == 'Text 30':
            set_shape_text(sh, '04', Pt(14), True)
        elif name == 'Text 31':
            set_shape_text(sh, 'Chuẩn hóa văn phong & Ngữ pháp', Pt(11), True)
        elif name == 'Text 32':
            set_shape_text(sh, 'Rà soát chính tả tiếng Việt, tinh chỉnh câu từ mạch lạc và chuẩn mực theo văn bản hành chính Công đoàn.', Pt(9.5))
        elif name == 'Text 34':
            set_shape_text(sh, 'Cơ sở kỹ thuật: Tích hợp Google Gemini API qua giao thức an toàn, xử lý dữ liệu thời gian thực và quản lý hạn mức Token.', Pt(9.5))

# ==================== SLIDE 7 ====================
s7 = prs.slides[6]
for sh in s7.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            set_shape_text(sh, 'Kiến trúc kỹ thuật — 4 thành phần hệ thống', Pt(18), True)
        elif name == 'Text 7': set_shape_text(sh, 'Frontend', Pt(11), True)
        elif name == 'Text 8': set_shape_text(sh, 'GIAO DIỆN NGƯỜI DÙNG', Pt(9), True)
        elif name == 'Text 11': set_shape_text(sh, 'React 18 SPA · Tailwind CSS', Pt(10.5), True)
        elif name == 'Text 14': set_shape_text(sh, 'Cổng thông tin và CMS cơ chế tải nhanh Single Page Application, chuẩn nhận diện TDMU, tối ưu hiển thị Responsive trên PC và di động.', Pt(9))
        elif name == 'Text 20': set_shape_text(sh, 'Backend', Pt(11), True)
        elif name == 'Text 21': set_shape_text(sh, 'XỬ LÝ NGHIỆP VỤ', Pt(9), True)
        elif name == 'Text 22': set_shape_text(sh, 'Node.js · Express RESTful API', Pt(10.5), True)
        elif name == 'Text 24': set_shape_text(sh, 'Điều phối nghiệp vụ tập trung', Pt(9.5), True)
        elif name == 'Text 28': set_shape_text(sh, 'Bảo mật & Tác vụ ngầm', Pt(9.5), True)
        elif name == 'Text 29': set_shape_text(sh, 'Xử lý logic tập trung, xác thực bảo mật JWT, mã hóa Bcrypt và điều phối tiến trình ngầm (Node-cron) lập lịch xuất bản.', Pt(9))
        elif name == 'Text 35': set_shape_text(sh, 'Cơ sở dữ liệu', Pt(11), True)
        elif name == 'Text 36': set_shape_text(sh, 'LƯU TRỮ & TOÀN VẸN', Pt(9), True)
        elif name == 'Text 37': set_shape_text(sh, 'Microsoft SQL Server 2022', Pt(10.5), True)
        elif name == 'Text 42': set_shape_text(sh, 'Mô hình dữ liệu quan hệ chuẩn hóa 3NF, bảo đảm toàn vẹn giao dịch ACID, tối ưu hóa truy vấn thông qua hệ thống chỉ mục (Indexing).', Pt(9))
        elif name == 'Text 48': set_shape_text(sh, 'Dịch vụ tích hợp', Pt(11), True)
        elif name == 'Text 49': set_shape_text(sh, 'AI & XUẤT BẢN ĐA KÊNH', Pt(9), True)
        elif name == 'Text 50': set_shape_text(sh, 'Google Gemini API', Pt(10.5), True)
        elif name == 'Text 53': set_shape_text(sh, 'Meta Graph API v19', Pt(10.5), True)
        elif name == 'Text 54': set_shape_text(sh, 'Gemini xử lý ngôn ngữ tự nhiên hỗ trợ tạo bài viết; Meta Graph API cung cấp kết nối an toàn tự động đăng bài lên Fanpage.', Pt(9))

# ==================== SLIDE 8 ====================
s8 = prs.slides[7]
for sh in s8.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            set_shape_text(sh, 'Biểu đồ Use Case tổng quát và 4 nhóm tác nhân', Pt(18), True)
        elif name == 'Text 2':
            set_shape_text(sh, '4 NHÓM TÁC NHÂN (ACTOR)', Pt(11), True)
        elif name == 'Text 3':
            set_shape_text(sh, 'Quản trị viên (Admin): Quản lý tài khoản, phân quyền, cấu hình hệ thống & Token AI', Pt(9.5))
        elif name == 'Text 4':
            set_shape_text(sh, 'Biên tập viên (Editor): Quản lý bài viết, biên tập, hẹn giờ và xuất bản đa kênh', Pt(9.5))
        elif name == 'Text 5':
            set_shape_text(sh, 'Cộng tác viên (Contributor): Soạn thảo bài viết, tải tư liệu sự kiện và gửi bài lên', Pt(9.5))
        elif name == 'Text 6':
            set_shape_text(sh, 'Đoàn viên / Độc giả (User): Đọc báo, tra cứu văn bản, biểu mẫu và tương tác', Pt(9.5))
        elif name == 'Text 7':
            set_shape_text(sh, '5 PHÂN HỆ NGHIỆP VỤ (26 USE CASE)', Pt(11), True)
        elif name == 'Text 11':
            set_shape_text(sh, '1. Quản lý bài viết & nội dung — Soạn thảo, đính kèm media, lưu bản nháp, phân quyền tác giả', Pt(9))
        elif name == 'Text 14':
            set_shape_text(sh, '2. Trợ lý AI hỗ trợ sáng tạo — Dựng bài 5W1H, gợi ý tiêu đề, tóm tắt Sapo, tối ưu đa kênh', Pt(9))
        elif name == 'Text 17':
            set_shape_text(sh, '3. Xuất bản đa kênh & Lập lịch — Xem trước định dạng, hẹn giờ đăng bài, đồng bộ Facebook', Pt(9))
        elif name == 'Text 23':
            set_shape_text(sh, '4. Cổng thông tin độc giả — Đọc tin tức, tra cứu văn bản pháp quy, tải biểu mẫu, gửi góp ý', Pt(9))
        elif name == 'Text 25':
            set_shape_text(sh, '5. Quản trị hệ thống & Người dùng — Quản lý tài khoản, phân quyền RBAC, cấu hình API, audit logs', Pt(9))
        elif name == 'Text 27':
            set_shape_text(sh, '26 use case được đặc tả chuẩn UML với đầy đủ quan hệ <<include>>, <<extend>> và luồng sự kiện chi tiết.', Pt(9))

# ==================== SLIDE 9 ====================
s9 = prs.slides[8]
for sh in s9.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 1':
            set_shape_text(sh, 'Thiết kế cơ sở dữ liệu quan hệ (15 bảng chuẩn 3NF)', Pt(18), True)
        elif name == 'Text 3':
            set_shape_text(sh, 'Đặc tả kỹ thuật CSDL', Pt(11), True)
        elif name == 'Text 4':
            set_shape_text(sh, '• Microsoft SQL Server 2022 kết hợp Connection Pool tối ưu hiệu năng kết nối.\n• Chuẩn hóa dữ liệu bậc 3 (3NF), loại bỏ dư thừa và hạn chế dị thường dữ liệu.\n• Ràng buộc toàn vẹn khóa chính (PK) / khóa ngoại (FK) và bảo toàn giao dịch ACID.\n• Đánh chỉ mục (Indexing) trên các trường slug, category_id, status tăng tốc truy vấn.', Pt(9.5))
        elif name == 'Text 6':
            set_shape_text(sh, 'Cấu trúc 15 bảng theo 4 nhóm thực thể chính:', Pt(10.5), True)

# ==================== SLIDE 10 ====================
s10 = prs.slides[9]
for sh in s10.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0': set_shape_text(sh, '5. KẾT QUẢ THỰC NGHIỆM', Pt(10), True)
        elif name == 'Text 1': set_shape_text(sh, 'Giao diện Cổng thông tin độc giả (User đọc báo & tra cứu)', Pt(18), True)
        elif name == 'Text 2': set_shape_text(sh, 'Chuẩn nhận diện TDMU', Pt(9.5), True)
        elif name == 'Text 4': set_shape_text(sh, 'Kho Văn bản & Biểu mẫu số hóa', Pt(9.5), True)
        elif name == 'Text 6': set_shape_text(sh, 'Đọc bài viết chuẩn báo chí', Pt(9.5), True)
        elif name == 'Text 8': set_shape_text(sh, 'Tối ưu Responsive đa thiết bị', Pt(9.5), True)
        elif name == 'Text 11': set_shape_text(sh, 'Trang chủ Cổng thông tin', Pt(11), True)
        elif name == 'Text 12': set_shape_text(sh, 'Bố cục trang nhã theo chuẩn nhận diện TDMU; menu chuyên mục rõ ràng, tích hợp banner sự kiện nổi bật và luồng tin phong trào mới nhất.', Pt(9))
        elif name == 'Text 14': set_shape_text(sh, 'Trang chi tiết bài viết & Tương tác', Pt(11), True)
        elif name == 'Text 15': set_shape_text(sh, 'Định dạng bài báo chuẩn mực (Tiêu đề, Sapo, Toàn văn); hiển thị album ảnh sự kiện, tệp văn bản đính kèm và khu vực phản hồi độc giả.', Pt(9))

# ==================== SLIDE 11 ====================
s11 = prs.slides[10]
for sh in s11.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0': set_shape_text(sh, '5. KẾT QUẢ THỰC NGHIỆM', Pt(10), True)
        elif name == 'Text 1': set_shape_text(sh, 'Phân hệ Quản trị CMS & Phòng biên tập tạo tin AI đa kênh', Pt(18), True)
        elif name == 'Text 2': set_shape_text(sh, 'Bảng điều khiển theo dõi tin bài', Pt(9.5), True)
        elif name == 'Text 4': set_shape_text(sh, 'Biên tập viên quản lý xuất bản', Pt(9.5), True)
        elif name == 'Text 6': set_shape_text(sh, 'Cộng tác viên tạo bài bằng AI', Pt(9.5), True)
        elif name == 'Text 8': set_shape_text(sh, 'Xuất bản đa kênh tự động', Pt(9.5), True)
        elif name == 'Text 11': set_shape_text(sh, 'Phòng Biên tập tích hợp Trợ lý AI', Pt(11), True)
        elif name == 'Text 12': set_shape_text(sh, 'Soạn thảo bài viết trực quan; Trợ lý AI hỗ trợ tạo bài, gợi ý tiêu đề và xem trước định dạng bài đăng Fanpage Facebook, Zalo.', Pt(9))
        elif name == 'Text 14': set_shape_text(sh, 'Bảng Điều Hành (Dashboard)', Pt(11), True)
        elif name == 'Text 15': set_shape_text(sh, 'Quản lý tập trung bài viết, lịch hẹn giờ đăng và theo dõi tổng thể số liệu tương tác đa kênh của hệ thống.', Pt(9))

# ==================== SLIDE 12 ====================
# Card 1: Text 16 (Title), Text 17 (Desc)
# Card 2: Text 24 (Title), Text 25 (Desc)
# Bottom Box: Text 36 (Title), Text 37 (Desc)
s12 = prs.slides[11]
for sh in s12.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 0': set_shape_text(sh, '6. ĐÁNH GIÁ & HƯỚNG PHÁT TRIỂN', Pt(10), True)
        elif name == 'Text 1': set_shape_text(sh, 'Kết quả đạt được và Hướng phát triển đề tài', Pt(18), True)
        elif name == 'Text 16': set_shape_text(sh, 'TẠO BÀI VIẾT BẰNG AI', Pt(11), True)
        elif name == 'Text 17': set_shape_text(sh, 'AI tự động phân tích tư liệu sự kiện, hỗ trợ lập dàn ý, viết bài 5W1H và chuẩn hóa câu từ, giảm tải thời gian biên soạn cho cán bộ.', Pt(9.5))
        elif name == 'Text 24': set_shape_text(sh, 'XUẤT BẢN ĐA KÊNH & HẸN GIỜ', Pt(11), True)
        elif name == 'Text 25': set_shape_text(sh, 'Tự động đăng tải bài viết và hình ảnh lên Fanpage Facebook ngay sau khi xuất bản; hỗ trợ lập lịch hẹn giờ phát hành ngầm chính xác.', Pt(9.5))
        elif name == 'Text 36': set_shape_text(sh, 'Hướng phát triển đề tài', Pt(11), True)
        elif name == 'Text 37': set_shape_text(sh, '• Tinh chỉnh và tối ưu hóa bộ prompt chuyên sâu cho từng loại văn bản và hoạt động công đoàn đặc thù.\n• Hoàn tất kết nối Webhook chính thức với Zalo Official Account khi được cấp phép.\n• Nâng cấp các tiện ích khảo sát và lắng nghe tâm tư nguyện vọng của đoàn viên cơ sở.\n• Triển khai lấy ý kiến thử nghiệm thực tế từ các Tổ công đoàn để tối ưu trải nghiệm người dùng.', Pt(9.5))

# ==================== SLIDE 13 ====================
s13 = prs.slides[12]
for sh in s13.shapes:
    if sh.has_text_frame:
        name = sh.name
        if name == 'Text 39':
            set_shape_text(sh, 'CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ LẮNG NGHE!', Pt(24), True)
        elif name == 'Hộp Văn bản 2':
            set_shape_text(sh, 'Nhóm 2 xin trân trọng cảm ơn sự hướng dẫn tận tình của ThS. Võ Quốc Lương!\nKính mời Thầy nhận xét và đặt câu hỏi cho nhóm em ạ!', Pt(14), False)

prs.save(dst)
print('Successfully saved thuyet trinh.pptx perfectly aligned with pre-existing layout!')
