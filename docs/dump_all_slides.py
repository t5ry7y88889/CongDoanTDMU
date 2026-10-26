import pptx, sys
sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation(r'Z:\Tailieu\DoAnCSNganh\tdmu-congdoan-web\docs\thuyet trinh.pptx')

for i in range(len(prs.slides)):
    print(f'=== SLIDE {i+1} ===')
    s = prs.slides[i]
    for idx, sh in enumerate(s.shapes):
        txt = sh.text_frame.text.replace('\n', ' // ') if sh.has_text_frame else ''
        if txt or sh.shape_type == pptx.enum.shapes.MSO_SHAPE.RECTANGLE:
            print(f'[{idx}] {sh.name} id={sh.shape_id} pos=({sh.left},{sh.top},{sh.width},{sh.height}) txt="{txt[:70]}"')
